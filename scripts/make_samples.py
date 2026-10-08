"""Create the small sample files used by the P3 engine spike and later tests.

Known strings live here so a spike can check that a tool read the file,
not that it invented a value.
"""

from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "sample_docs"

HEADING = "Annual Financial Report 2025"
BODY = "The company delivered strong financial performance during FY2025."
SCAN_TEXT = "Scanned revenue note total is 128.5 crore."
EQUATION = "E = mc^2"


def make_simple(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 80), HEADING, fontsize=18)
    page.insert_text((72, 120), BODY, fontsize=12)
    doc.save(path)
    doc.close()


def make_scanned(path: Path) -> None:
    """Image-only PDF. No text layer, so OCR is the only way to read it."""
    src = pymupdf.open()
    page = src.new_page(width=612, height=792)
    page.insert_text((72, 120), SCAN_TEXT, fontsize=16)
    pix = page.get_pixmap(dpi=150)

    out = pymupdf.open()
    img_page = out.new_page(width=612, height=792)
    img_page.insert_image(img_page.rect, pixmap=pix)
    out.save(path)
    out.close()
    src.close()


def make_table(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    headers = ["Year", "Revenue", "EBITDA"]
    rows = [["2024", "110.7", "23.5"], ["2025", "128.5", "29.1"]]
    x0, y0 = 72, 80
    col_w, row_h = 120, 28
    grid = [headers, *rows]
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(
                x0 + c * col_w,
                y0 + r * row_h,
                x0 + (c + 1) * col_w,
                y0 + (r + 1) * row_h,
            )
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            page.insert_text((rect.x0 + 8, rect.y0 + 18), value, fontsize=11)
    doc.save(path)
    doc.close()


def make_equation(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 80), "Equation sample", fontsize=14)
    page.insert_text((72, 140), EQUATION, fontsize=20)
    doc.save(path)
    doc.close()


def make_chart(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 60), "Figure 1 Revenue", fontsize=14)
    labels = [("2023", 92.4), ("2024", 110.7), ("2025", 128.5)]
    x = 90
    for label, value in labels:
        height = value * 2
        rect = pymupdf.Rect(x, 400 - height, x + 50, 400)
        page.draw_rect(rect, color=(0.1, 0.3, 0.6), fill=(0.2, 0.4, 0.7))
        page.insert_text((x, 418), label, fontsize=11)
        page.insert_text((x, 400 - height - 14), str(value), fontsize=11)
        x += 90
    doc.save(path)
    doc.close()


def make_xlsx(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Financials"
    sheet["A1"] = "Year"
    sheet["B1"] = "Revenue"
    sheet["A2"] = 2025
    sheet["B2"] = 128.5
    book.save(path)


def make_docx(path: Path) -> None:
    document = Document()
    document.add_heading(HEADING, level=1)
    document.add_paragraph(BODY)
    document.save(path)


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Revenue Growth"
    chart_data = CategoryChartData()
    chart_data.categories = ["2023", "2024", "2025"]
    chart_data.add_series("Revenue", (92.4, 110.7, 128.5))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(2),
        Inches(6),
        Inches(4),
        chart_data,
    )
    presentation.save(path)


def _draw_grid(page, rows: list[list[str]], x0: float, y0: float) -> None:
    col_w, row_h = 110, 28
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(
                x0 + c * col_w,
                y0 + r * row_h,
                x0 + (c + 1) * col_w,
                y0 + (r + 1) * row_h,
            )
            page.draw_rect(rect, color=(0, 0, 0), width=0.6)
            page.insert_text((rect.x0 + 6, rect.y0 + 18), value, fontsize=10)


def make_ruled_cross_page_table(path: Path) -> None:
    """Two ruled tables. Page 1 ends near the bottom. Page 2 starts near the top."""
    header = ["Company", "Revenue", "EBITDA", "Margin"]
    page1_rows = [
        header,
        ["Alpha", "128.5", "32.1", "25.0%"],
        ["Beta", "94.2", "21.7", "23.0%"],
        ["Gamma", "76.8", "18.9", "24.6%"],
        ["Delta", "61.4", "12.6", "20.5%"],
        ["Epsilon", "58.7", "11.4", "19.4%"],
    ]
    page2_rows = [
        header,
        ["Zeta", "52.3", "10.8", "20.7%"],
        ["Eta", "49.6", "13.2", "26.6%"],
        ["Theta", "43.1", "8.7", "20.2%"],
        ["Iota", "39.8", "7.9", "19.8%"],
        ["Kappa", "34.5", "6.2", "18.0%"],
    ]
    doc = pymupdf.open()
    first = doc.new_page(width=612, height=792)
    first.insert_text((72, 80), "Operating Results", fontsize=16)
    first.insert_text((72, 110), "This ruled table continues on the next page.", fontsize=11)
    _draw_grid(first, page1_rows, 72, 600)
    second = doc.new_page(width=612, height=792)
    second.insert_text((72, 250), "Continued from the previous page.", fontsize=11)
    _draw_grid(second, page2_rows, 72, 36)
    doc.save(path)
    doc.close()


def make_empty_pdf(path: Path) -> None:
    path.write_bytes(b"")


def make_unsupported(path: Path) -> None:
    path.write_text(
        "This file is plain text. Preflight should reject the extension.\n",
        encoding="utf-8",
    )


def main() -> None:
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    make_simple(SAMPLE_DIR / "simple.pdf")
    make_scanned(SAMPLE_DIR / "scanned.pdf")
    make_table(SAMPLE_DIR / "table.pdf")
    make_equation(SAMPLE_DIR / "equation.pdf")
    make_chart(SAMPLE_DIR / "chart.pdf")
    make_xlsx(SAMPLE_DIR / "sample.xlsx")
    make_docx(SAMPLE_DIR / "sample.docx")
    make_pptx(SAMPLE_DIR / "sample.pptx")
    make_ruled_cross_page_table(SAMPLE_DIR / "04_ruled_cross_page_table.pdf")
    make_empty_pdf(SAMPLE_DIR / "11_empty.pdf")
    make_unsupported(SAMPLE_DIR / "11_notes.txt")
    print(f"Wrote samples in {SAMPLE_DIR}")


if __name__ == "__main__":
    main()
