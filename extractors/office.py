"""Native Office extraction. No vision model.

P1 calls extract_document(path) once per file. route="office" does the same.
"""

from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from extractors.base import Extractor
from extractors.tables import assemble
from extractors.utils import failed_block, make_block, region_field

_EMPTY_BOX = [0.0, 0.0, 0.0, 0.0]


class OfficeExtractor(Extractor):
    name = "office"

    def can_handle(self, region) -> bool:
        return region_field(region, "route") == "office"

    def extract(self, region, context) -> list:
        context = context or {}
        path = context.get("file_path")
        if not path:
            return [_fail(f"Office extraction needs a file path")]
        return extract_document(path)


def extract_document(path) -> list:
    """Read one DOCX, XLSX, or PPTX into blocks. This is the function P1 calls."""
    file_path = Path(path)
    if not file_path.is_file():
        return [_fail(f"Office file not found: {path}")]
    kind = file_path.suffix.lower()
    if kind == ".xlsx":
        return _xlsx(file_path)
    if kind == ".docx":
        return _docx(file_path)
    if kind == ".pptx":
        return _pptx(file_path)
    return [_fail(f"Unsupported office format: {kind}")]


def _xlsx(path: Path) -> list:
    book = load_workbook(path, data_only=False)
    blocks = []
    try:
        for sheet in book.worksheets:
            cells = _sheet_cells(sheet)
            if not cells:
                continue
            content = assemble(cells)
            span = sheet.dimensions or _range_from_cells(cells)
            blocks.append(
                make_block(
                    extractor="openpyxl",
                    bbox=_EMPTY_BOX,
                    page_start=0,
                    source_page=None,
                    extraction=1.0,
                    block_type="table",
                    content=content,
                    region_id=path.stem,
                    index=len(blocks),
                    source_file=str(path),
                    sheet=sheet.title,
                    cell_range=span,
                    history=[{"engine": "openpyxl", "confidence": 1.0, "note": sheet.title}],
                )
            )
    finally:
        book.close()
    return blocks or [_fail(f"No values in {path.name}")]


def _sheet_cells(sheet) -> list[dict]:
    covered = set()
    spans = {}
    for merged in sheet.merged_cells.ranges:
        spans[(merged.min_row, merged.min_col)] = (
            merged.max_row - merged.min_row + 1,
            merged.max_col - merged.min_col + 1,
        )
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if (row, col) != (merged.min_row, merged.min_col):
                    covered.add((row, col))
    cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None or (cell.row, cell.column) in covered:
                continue
            row_span, col_span = spans.get((cell.row, cell.column), (1, 1))
            formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
            text = "" if formula else str(cell.value)
            cells.append(
                {
                    "row": cell.row - 1,
                    "col": cell.column - 1,
                    "row_span": row_span,
                    "col_span": col_span,
                    "text": text,
                    "bbox": _EMPTY_BOX,
                    "confidence": 1.0,
                    "formula": formula,
                    "number_format": cell.number_format,
                }
            )
    return cells


def _docx(path: Path) -> list:
    document = Document(str(path))
    blocks = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.startswith("Heading"):
            level_text = style.replace("Heading", "").strip()
            level = int(level_text) if level_text.isdigit() else 1
            content = {"text": text, "level": level}
            block_type = "heading"
        elif _is_list(paragraph):
            content = {"text": text}
            block_type = "list"
        else:
            content = {"text": text}
            block_type = "paragraph"
        blocks.append(_office_block(path, "python-docx", block_type, content, len(blocks)))
    for table in document.tables:
        cells = []
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                cells.append(
                    {
                        "row": row_index,
                        "col": col_index,
                        "row_span": 1,
                        "col_span": 1,
                        "text": cell.text.strip(),
                        "bbox": _EMPTY_BOX,
                        "confidence": 1.0,
                    }
                )
        if cells:
            blocks.append(
                _office_block(path, "python-docx", "table", assemble(cells), len(blocks))
            )
    return blocks or [_fail(f"No content in {path.name}")]


def _pptx(path: Path) -> list:
    presentation = Presentation(str(path))
    blocks = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        for shape in slide.shapes:
            bbox = _shape_bbox(shape)
            if shape.has_text_frame:
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
                if text:
                    is_title = title_shape is not None and shape == title_shape
                    content = {"text": text, "level": 1} if is_title else {"text": text}
                    blocks.append(
                        _office_block(
                            path,
                            "python-pptx",
                            "heading" if is_title else "paragraph",
                            content,
                            len(blocks),
                            bbox=bbox,
                            slide=slide_number,
                        )
                    )
            if shape.has_table:
                cells = _pptx_table_cells(shape.table)
                if cells:
                    blocks.append(
                        _office_block(
                            path,
                            "python-pptx",
                            "table",
                            assemble(cells),
                            len(blocks),
                            bbox=bbox,
                            slide=slide_number,
                        )
                    )
            if shape.has_chart:
                blocks.append(
                    _office_block(
                        path,
                        "python-pptx",
                        "chart",
                        _chart_content(shape.chart),
                        len(blocks),
                        bbox=bbox,
                        slide=slide_number,
                    )
                )
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blocks.append(
                    _office_block(
                        path,
                        "python-pptx",
                        "figure",
                        {"text": shape.name},
                        len(blocks),
                        bbox=bbox,
                        slide=slide_number,
                    )
                )
    return blocks or [_fail(f"No content in {path.name}")]


def _chart_content(chart) -> dict:
    ticks = _categories(chart)
    series = []
    for item in chart.series:
        values = [None if value is None else float(value) for value in item.values]
        points = [
            {"label": ticks[index] if index < len(ticks) else str(index + 1), "value": value}
            for index, value in enumerate(values)
        ]
        series.append({"name": str(item.name), "points": points, "method": "pptx_native"})
    kind = str(chart.chart_type)
    if "LINE" in kind:
        chart_type = "line"
    elif "COLUMN" in kind or "BAR" in kind:
        chart_type = "bar"
    else:
        chart_type = kind
    return {
        "chart_type": chart_type,
        "title": None,
        "x_axis": {"label": None, "ticks": ticks},
        "y_axis": {"label": series[0]["name"] if series else None, "unit": None},
        "series": series,
        "agreement": None,
        "candidates": [],
    }


def _categories(chart) -> list[str]:
    try:
        return [str(category) for category in chart.plots[0].categories]
    except Exception:
        return []


def _pptx_table_cells(table) -> list[dict]:
    cells = []
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cells.append(
                {
                    "row": row_index,
                    "col": col_index,
                    "row_span": 1,
                    "col_span": 1,
                    "text": cell.text.strip(),
                    "bbox": _EMPTY_BOX,
                    "confidence": 1.0,
                }
            )
    return cells


def _office_block(path, extractor, block_type, content, index, bbox=None, slide=None):
    box = bbox or _EMPTY_BOX
    return make_block(
        extractor=extractor,
        bbox=box,
        page_start=slide or 0,
        source_page=slide,
        extraction=1.0,
        block_type=block_type,
        content=content,
        region_id=path.stem,
        index=index,
        source_file=str(path),
        slide=slide,
        history=[{"engine": extractor, "confidence": 1.0, "note": block_type}],
    )


def _shape_bbox(shape) -> list[float]:
    left = float(shape.left)
    top = float(shape.top)
    return [
        left / 12700,
        top / 12700,
        (left + float(shape.width)) / 12700,
        (top + float(shape.height)) / 12700,
    ]


def _is_list(paragraph) -> bool:
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        return True
    style = paragraph.style.name if paragraph.style is not None else ""
    return style.startswith("List")


def _range_from_cells(cells: list[dict]) -> str:
    rows = [cell["row"] + 1 for cell in cells]
    cols = [cell["col"] + 1 for cell in cells]
    start = f"{get_column_letter(min(cols))}{min(rows)}"
    end = f"{get_column_letter(max(cols))}{max(rows)}"
    return start if start == end else f"{start}:{end}"


def _fail(message: str):
    return failed_block(
        {"id": "office", "page": 0, "bbox": _EMPTY_BOX},
        error_code="INTERNAL_ERROR",
        message=message,
        extractor="office",
    )
